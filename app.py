from flask import Flask, render_template, request, jsonify, redirect, url_for, flash
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
import os
from estrai import EstrattorePDF
import logging
from logging.handlers import RotatingFileHandler

def setup_logging():
    if not os.path.exists('instance'):
        os.makedirs('instance')
    
    handler = RotatingFileHandler(
        'instance/api.log', 
        maxBytes=1024*1024,
        backupCount=5
    )
    
    formatter = logging.Formatter(
        '[%(asctime)s] [%(levelname)s] %(message)s',
        datefmt='%Y-%m-%d %H:%M:%S'
    )
    handler.setFormatter(formatter)
    
    logger = logging.getLogger('datacleaner_api')
    logger.setLevel(logging.INFO)
    logger.addHandler(handler)
    
    return logger

api_logger = setup_logging()

app = Flask(__name__)
app.config['SECRET_KEY'] = 'chiave-segreta-temporanea'
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///utenti.db'
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'

estrattore = EstrattorePDF()

class Utente(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(200), nullable=False)
    crediti = db.Column(db.Integer, default=10)  # 10 elaborazioni gratis

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

@login_manager.user_loader
def load_user(user_id):
    return Utente.query.get(int(user_id))

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/registrati', methods=['GET', 'POST'])
def registrati():
    if request.method == 'POST':
        email = request.form.get('email')
        password = request.form.get('password')
        
        if Utente.query.filter_by(email=email).first():
            return jsonify({'errore': 'Email già registrata'}), 400
        
        utente = Utente(email=email)
        utente.set_password(password)
        db.session.add(utente)
        db.session.commit()
        
        login_user(utente)
        return jsonify({'successo': True, 'crediti': utente.crediti})
    
    return render_template('registrati.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email')
        password = request.form.get('password')
        
        utente = Utente.query.filter_by(email=email).first()
        
        if utente and utente.check_password(password):
            login_user(utente)
            return jsonify({'successo': True, 'crediti': utente.crediti})
        
        return jsonify({'errore': 'Credenziali non valide'}), 401
    
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('home'))

@app.route('/crediti')
@login_required
def crediti():
    return jsonify({'crediti': current_user.crediti})

@app.route('/elabora', methods=['POST'])
@login_required
def elabora():
    if current_user.crediti <= 0:
        return jsonify({'errore': 'Crediti esauriti. Acquista nuovi crediti.'}), 403
    
    if 'file' not in request.files:
        return jsonify({'errore': 'Nessun file caricato'}), 400
    
    file = request.files['file']
    
    if file.filename == '':
        return jsonify({'errore': 'Nessun file selezionato'}), 400
    
    percorso_temp = 'temp.pdf'
    file.save(percorso_temp)
    
    risultato = estrattore.processa_pdf(percorso_temp)
    
    if os.path.exists(percorso_temp):
        os.remove(percorso_temp)
    
    if not risultato:
        return jsonify({'errore': 'Impossibile elaborare il PDF'}), 400
    
    current_user.crediti -= 1
    db.session.commit()
    
    risultato['crediti_rimasti'] = current_user.crediti
    
    return jsonify(risultato)

@app.route('/scarica-excel', methods=['POST'])
@login_required
def scarica_excel():
    dati = request.get_json()
    movimenti = dati.get('movimenti', [])
    
    if not movimenti:
        return jsonify({'errore': 'Nessun dato da scaricare'}), 400
    
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimenti"
    
    headers = ['Data', 'Descrizione', 'Entrate', 'Uscite', 'Saldo']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4ecca3", end_color="4ecca3", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    for row, mov in enumerate(movimenti, 2):
        ws.cell(row=row, column=1, value=mov.get('data', ''))
        ws.cell(row=row, column=2, value=mov.get('descrizione', ''))
        ws.cell(row=row, column=3, value=mov.get('entrata') if mov.get('entrata') else None)
        ws.cell(row=row, column=4, value=mov.get('uscita') if mov.get('uscita') else None)
        ws.cell(row=row, column=5, value=mov.get('saldo') if mov.get('saldo') else None)
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    from flask import Response
    return Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=movimenti.xlsx'}
    )


@app.route('/scarica-csv', methods=['POST'])
@login_required
def scarica_csv():
    dati = request.get_json()
    movimenti = dati.get('movimenti', [])
    
    if not movimenti:
        return jsonify({'errore': 'Nessun dato da scaricare'}), 400
    
    import csv
    import io
    
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['Data', 'Descrizione', 'Entrate', 'Uscite', 'Saldo'])
    
    for mov in movimenti:
        writer.writerow([
            mov.get('data', ''),
            mov.get('descrizione', ''),
            mov.get('entrata', ''),
            mov.get('uscita', ''),
            mov.get('saldo', '')
        ])
    
    from flask import Response
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=dati_estratti.csv'}
    )


# Rate limiting semplice in memoria
richieste_per_ip = {}
LIMITE_ORARIO = 5
FINESTRA_SECONDI = 3600

def check_rate_limit(ip):
    import time
    ora_attuale = time.time()
    
    if ip not in richieste_per_ip:
        richieste_per_ip[ip] = []
    
    # Pulisci richieste vecchie
    richieste_per_ip[ip] = [
        t for t in richieste_per_ip[ip] 
        if ora_attuale - t < FINESTRA_SECONDI
    ]
    
    if len(richieste_per_ip[ip]) >= LIMITE_ORARIO:
        return False
    
    richieste_per_ip[ip].append(ora_attuale)
    return True


@app.route('/api/v1/converti', methods=['POST'])
def api_converti():
    # RATE LIMITING
    ip_cliente = request.remote_addr
    if not check_rate_limit(ip_cliente):
        api_logger.warning(f"RATE LIMIT | IP: {ip_cliente}")
        return jsonify({
            "stato": "errore", 
            "messaggio": "Rate limit superato. Riprova tra un'ora."
        }), 429
    
    # VALIDAZIONE 1: file presente
    if 'file' not in request.files:
        api_logger.warning(f"NO FILE | IP: {ip_cliente}")
        return jsonify({"stato": "errore", "messaggio": "Nessun file inviato"}), 400
    
    file_caricato = request.files['file']
    
    # VALIDAZIONE 2: nome file valido
    if file_caricato.filename == '':
        api_logger.warning(f"NOME VUOTO | IP: {ip_cliente}")
        return jsonify({"stato": "errore", "messaggio": "Nome file vuoto"}), 400
    
    # VALIDAZIONE 3: nome file sicuro
    nome_file_originale = file_caricato.filename
    nome_file = nome_file_originale.lower()
    
    if '/' in nome_file_originale or '\\' in nome_file_originale or '..' in nome_file_originale:
        api_logger.warning(f"PATH TRAVERSAL | IP: {ip_cliente} | File: {nome_file_originale}")
        return jsonify({"stato": "errore", "messaggio": "Nome file non valido"}), 400
    
    # VALIDAZIONE 4: estensione supportata
    if not (nome_file.endswith('.csv') or nome_file.endswith('.xlsx') or nome_file.endswith('.xls')):
        api_logger.warning(f"FORMATO NON SUPPORTATO | IP: {ip_cliente} | File: {nome_file_originale}")
        return jsonify({"stato": "errore", "messaggio": "Formato non supportato. Usa CSV o Excel"}), 400
    
    # VALIDAZIONE 5: file non vuoto
    file_caricato.seek(0, 2)
    dimensione = file_caricato.tell()
    file_caricato.seek(0)
    
    if dimensione == 0:
        api_logger.warning(f"FILE VUOTO | IP: {ip_cliente} | File: {nome_file_originale}")
        return jsonify({"stato": "errore", "messaggio": "Il file è vuoto"}), 400
    
    # VALIDAZIONE 6: file non troppo piccolo
    if dimensione < 10:
        api_logger.warning(f"FILE TROPPO PICCOLO | IP: {ip_cliente} | File: {nome_file_originale} | Dim: {dimensione}")
        return jsonify({"stato": "errore", "messaggio": "Il file è troppo piccolo per essere valido"}), 400
    
    # Salvataggio temporaneo
    import time
    timestamp = str(int(time.time()))
    
    if nome_file.endswith('.csv'):
        percorso_temp = os.path.join('instance', 'temp_converti_' + timestamp + '.csv')
    elif nome_file.endswith('.xlsx'):
        percorso_temp = os.path.join('instance', 'temp_converti_' + timestamp + '.xlsx')
    else:
        percorso_temp = os.path.join('instance', 'temp_converti_' + timestamp + '.xls')
    
    # LOG INIZIO ELABORAZIONE
    api_logger.info(f"RICHIESTA | IP: {ip_cliente} | File: {nome_file_originale} | Dim: {dimensione} byte")
    
    try:
        file_caricato.save(percorso_temp)
        
        from trasforma_dati import normalizza_e_converti_foglio
        risultato = normalizza_e_converti_foglio(percorso_temp)
        
        if os.path.exists(percorso_temp):
            os.remove(percorso_temp)
        
        # VALIDAZIONE 7: risultato valido
        if risultato.get('stato') == 'errore':
            api_logger.error(f"ERRORE CONVERSIONE | IP: {ip_cliente} | File: {nome_file_originale}")
            return jsonify(risultato), 422
        
        # VALIDAZIONE 8: almeno una riga
        righe = risultato.get('info_file', {}).get('numero_righe_estratte', 0)
        if righe == 0:
            api_logger.warning(f"NESSUN DATO | IP: {ip_cliente} | File: {nome_file_originale}")
            return jsonify({
                "stato": "errore", 
                "messaggio": "Il file non contiene righe valide dopo la pulizia"
            }), 422
        
        # LOG SUCCESSO
        api_logger.info(f"SUCCESSO | IP: {ip_cliente} | File: {nome_file_originale} | Righe: {righe}")
        
        return jsonify(risultato), 200
        
    except Exception as e:
        if os.path.exists(percorso_temp):
            os.remove(percorso_temp)
        
        api_logger.error(f"ECCEZIONE | IP: {ip_cliente} | Errore: {str(e)}")
        
        return jsonify({
            "stato": "errore", 
            "messaggio": f"Errore durante l'elaborazione: {str(e)}"
        }), 500
if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True, host='0.0.0.0', port=5000)